# coding=utf-8
import docx
import xlsxwriter


workbook = xlsxwriter.Workbook(input('请输入正确的输出表格文件路径，保存格式建议为较稳定的xlsx：'))
# 创建输出表格文件
worksheet = workbook.add_worksheet()


if __name__ == '__main__':
    file = docx.Document(input('请输入要读取的正确的docx文件路径:'))

    # 创建字典储存标题值
    dict0 = {'序号': None, '年龄': None, '性别': None, '编号': None, '内外向(E)': None \
        , '神经质(N)': None, '精神质(P)': None, '掩饰性(L)': None, \
             '躯体化': None, '强迫症状': None, '人际关系敏感': None, '抑郁': None, '焦虑': None, '敌对': None, '恐怖': None, \
             '偏执': None, '精神病性': None, '其他': None, '总分': None, '总均分': None, '阳性项目数': None}

    list0 = [dict0]
    tables = file.tables
    for table in tables:
        if table.cell(0, 0).text == '编号':
            if int(table.cell(0, 3).text) == len(list0):
                a = int(table.cell(0, 3).text)
                list0.append(dict0.copy())
            for n in range(1, len(table.rows[0].cells), 2):
                list0[int(table.cell(0, 3).text)][table.cell(0, n - 1).text] = table.cell(0, n).text
        if table.cell(0, 0).text == '因子名称' and table.cell(1, 0).text in ['内外向(E)', '躯体化']:
            for m in range(1, len(table.rows)):
                list0[a][table.cell(m, 0).text] = table.cell(m, 2).text

a = 0;
b = 0

# 写入标题行
for key in dict0.keys():
    worksheet.write(a, b, key)
    b += 1
b = 0
a += 1
for n in range(1, len(list0)):
    for value in list0[n].values():
        worksheet.write(a, b, value)
        b += 1
    b = 0
    a += 1
workbook.close()
