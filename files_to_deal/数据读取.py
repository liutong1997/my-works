# coding=utf-8
import docx
from xlsxwriter import *

workbook = Workbook(input('请输入正确的输出表格文件路径，保存格式建议为较稳定的xlsx：'))
# 创建输出表格文件
worksheet = workbook.add_worksheet()


# 定义清空字典值而不改变键的函数
def value_clear(dict2={'1': 'a', "2": 's'}):
    key = dict2.keys()
    for i in key:
        dict2[i] = None
    return dict2


# 定义判断字符串是否只由浮点数或整数构成的函数
def strs_are_nums(a='22.22'):
    n = 0
    for i in range(len(a)):
        if '0' <= a[i] <= '9':
            n += 1
        if a[i] == '.' and '0' <= a[i + 1] <= '9' and '0' <= a[i - 1] <= '9':
            n += 1
    if n == len(a):
        return True
    else:
        return False


if __name__ == '__main__':

    file = docx.Document(input('请输入要读取的正确的docx文件路径:'))

    # 创建字典储存标题值
    dict0 = {'序号': None, '年龄': None, '性别': None, '编号': None, '内外向(E)': None \
        , '神经质(N)': None, '精神质(P)': None, '掩饰性(L)': None, \
             '躯体化': None, '强迫症状': None, '人际关系敏感': None, '抑郁': None, '焦虑': None, '敌对': None, '恐怖': None, \
             '偏执': None, '精神病性': None, '其他': None, '总分': None, '总均分': None, '阳性项目数': None}

    number = {}
    # k,a,b分别为用户数量，写入文件行数，列数
    k = 0;
    a = 0;
    b = 0

    # 写入标题行
    for key in dict0.keys():
        worksheet.write(a, b, key)
        b += 1
    a += 1

    tables = file.tables
    for table in tables:
        b = 0
        # 判断用户
        if table.cell(0, 1).text not in number and strs_are_nums(table.cell(0, 1).text):
            # 写入上一个用户信息
            if k > 1:
                print(dict0)
                for value in dict0.values():
                    worksheet.write(a, b, value)
                    b += 1
                a += 1
            # 利用字典储存用户编号，并且统计用户数量

            value_clear(dict0)
            number[table.cell(0, 1).text] = k
            k += 1
            # 遍历行
            for m in range(len(table.rows)):
                # 遍历列
                for n in range(len(table.rows[0].cells)):
                    # 判断是否为属性值
                    if table.cell(m, n).text in dict0:
                        if table.cell(m, n).text == '年龄':
                            print(table.cell(m, n+1).text)
                            dict0[table.cell(m, n).text] = table.cell(m, n + 1).text
                        # 如果符合性别值或者数字组成的字符串，则将值写入字典
                        if table.cell(m, n+1).text in ['男', '女'] or strs_are_nums(table.cell(m, n + 1).text):
                            dict0[table.cell(m, n).text] = table.cell(m, n + 1).text
                        # 利用try模块防止列表越界的情况，即n已经代表最后一列时，n+
                        try:
                            # 判断人格因子的参数，因为其原始分在第二列
                            if strs_are_nums(table.cell(m, n + 2).text):
                                dict0[table.cell(m, n).text] = table.cell(m, n + 2).text
                        # 直接跳过越界
                        except:
                            pass
        else:
            # 遍历行
            for m in range(len(table.rows)):
                # 遍历列
                for n in range(len(table.rows[0].cells)):
                    # 判断是否为属性值
                    if table.cell(m, n).text in dict0:
                        if table.cell(m, n).text == '年龄':
                            print(table.cell(m,n+1).text)
                            dict0[table.cell(m, n).text] = table.cell(m, n + 1).text
                        # 如果符合性别值或者数字组成的字符串，则将值写入字典
                        if table.cell(m, n + 1).text in ['男', '女'] or strs_are_nums(table.cell(m, n + 1).text):
                            dict0[table.cell(m, n).text] = table.cell(m, n + 1).text
                        # 利用try模块防止列表越界的情况，即n已经代表最后一列时，n+2是不存在的
                        try:
                            # 判断人格因子的参数，因为其原始分在第二列
                            if strs_are_nums(table.cell(m, n + 2).text):
                                dict0[table.cell(m, n).text] = table.cell(m, n + 2).text
                        # 直接跳过越界
                        except:
                            pass
    print(dict0)
    for value in dict0.values():
        worksheet.write(a, b, value)
        b += 1
    a += 1
    workbook.close()
